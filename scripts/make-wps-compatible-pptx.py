#!/usr/bin/env python3
"""Repackage the deck with one full-slide PNG per page for WPS compatibility."""

from __future__ import annotations

import argparse
import io
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def remove_default_slide(prs: Presentation) -> None:
    for slide in list(prs.slides):
        slide_id = slide._element.getparent()
        slide_id.remove(slide._element)


def build_compatible_deck(source: Path, output: Path) -> int:
    prs = Presentation()
    remove_default_slide(prs)
    prs.slide_width = Inches(13.3333333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    with zipfile.ZipFile(source) as archive:
        slide_count = len(
            [name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
        )
        if slide_count == 0:
            raise ValueError("Source deck has no slides")

        for index in range(1, slide_count + 1):
            media_name = f"ppt/media/image{index}.png"
            try:
                png = archive.read(media_name)
            except KeyError as exc:
                raise ValueError(f"Missing PNG fallback for slide {index}: {media_name}") from exc

            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                io.BytesIO(png),
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return slide_count


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    count = build_compatible_deck(args.source.resolve(), args.output.resolve())
    print(f"Created {args.output} with {count} raster-only slides")


if __name__ == "__main__":
    main()
